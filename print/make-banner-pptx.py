# -*- coding: utf-8 -*-
"""Builds print/banner.pptx: the wall banner as an editable PowerPoint.

One A4 landscape slide carrying both strips, German above and English below,
so the whole thing is a single sheet through the printer and one cut with a
guillotine. Every piece is a real shape and can be dragged, retyped or
recoloured.

Fonts are deliberately ones that ship with Office. The site's own faces
(Bitter, Cabin, Geist) are web fonts, not installed on a normal machine, so
naming them here would only produce a silent substitution that looks wrong.
Georgia stands in for Bitter and Trebuchet MS for Cabin.

Run:  python print/make-banner-pptx.py
"""
import os
import pathlib

import segno
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
# Two codes, each through the pub's own short redirect so a changed Google
# link or Instagram handle needs no reprint. The PNGs are regenerated on every
# run, so they cannot drift from these addresses.
CODES = [
    ("https://oberrieden.pub/review", ROOT / "print" / "qr-review.png"),
    ("https://oberrieden.pub/instagram", ROOT / "print" / "qr-instagram.png"),
]

MAROON = RGBColor(0x53, 0x1C, 0x23)
GOLD = RGBColor(0xDD, 0xBA, 0x7A)
PAPER = RGBColor(0xFB, 0xF7, 0xF0)
TAN = RGBColor(0x8A, 0x62, 0x39)

PAGE_W, PAGE_H = 297, 210   # A4 landscape
SHEET = 8                   # a printer cannot reach the paper edge
W, H = 281, 92              # the finished sign, in mm
BAR = 12                    # fascia board
MARGIN = 14
QRSIZE = 42                 # two codes; at 50 only one fits beside the text
QRGAP = 6
QRX = W - MARGIN - QRSIZE * 2 - QRGAP   # left edge of the first code
# 8 + 92 = 100, cut at 105, second strip 110 to 202, 8mm left at the foot.
STRIP_Y = (SHEET, 110)
CUT_Y = 105

CONTENT = [
    dict(quip="Neue Website. Immer noch das gleiche gute Bier.",
         ask="Hat es Ihnen gefallen? Sagen Sie es Google.\vWenn nicht, dem Wirt.",
         caps=("BEWERTEN SIE UNS", "FOLGEN SIE UNS")),
    dict(quip="New website. Still the same great beer.",
         ask="Enjoyed it? Tell Google.\vIf not, tell the landlord.",
         caps=("REVIEW US", "FOLLOW US")),
]

NOTES = (
    "One A4 landscape sheet, two strips, one cut. Finished size 281 x 92 mm each. Print at 100% scale with background graphics "
    "switched on, or the maroon bar will not appear.\n\n"
    "The left code points at oberrieden.pub/review, which redirects to the "
    "pub's Google listing; the right one at oberrieden.pub/instagram. If "
    "either link ever changes, the redirect is what changes - the codes stay "
    "valid, so printed copies keep working.\n\n"
    "Do not put this up until the website is live and approved."
)


def text(slide, x, y, w, h, body, *, font, size, colour,
         bold=False, spacing=None, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    # PowerPoint pads text boxes by default, which would push everything a
    # couple of millimetres off the alignment the shapes establish.
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    # A "\v" in the body is a deliberate line break inside the paragraph. A run
    # cannot carry one, so each line is its own run with a break between.
    for i, line in enumerate(body.split("\v")):
        if i:
            p.add_line_break()
        run = p.add_run()
        run.text = line
        f = run.font
        f.name, f.size, f.bold = font, Pt(size), bold
        f.color.rgb = colour
        if spacing is not None:
            # python-pptx exposes no letter spacing; spc is in hundredths of a point.
            run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    return box


def rect(slide, x, y, w, h, fill, line=None, line_pt=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y), Mm(w), Mm(h))
    shape.shadow.inherit = False          # python-pptx defaults to a drop shadow
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_pt)
    shape.text_frame.text = ""
    return shape


def strip(slide, oy, c):
    """One finished sign, drawn at x=SHEET, y=oy."""
    ox = SHEET

    # The edge of the sign, so it reads as an object hung on a dark wall,
    # and so the cut line is obvious.
    rect(slide, ox, oy, W, H, PAPER, line=MAROON, line_pt=1)
    # The fascia board.
    rect(slide, ox, oy, W, BAR, MAROON)

    text(slide, ox + MARGIN, oy, W - MARGIN * 2, BAR,
         "BIG BEN PUB  ·  OBERRIEDEN DORF",
         font="Trebuchet MS", size=11.5, colour=GOLD, bold=True,
         spacing=3, anchor=MSO_ANCHOR.MIDDLE)

    text(slide, ox + MARGIN, oy + 24, QRX - MARGIN - 6, 20, "oberrieden.pub",
         font="Georgia", size=50, colour=MAROON, bold=True)

    text(slide, ox + MARGIN, oy + 45.5, QRX - MARGIN - 6, 8, c["quip"],
         font="Georgia", size=18, colour=TAN, bold=True)

    text(slide, ox + MARGIN, oy + 55, QRX - MARGIN - 6, 13, c["ask"],
         font="Calibri", size=14, colour=MAROON)

    for i, ((_, png), cap) in enumerate(zip(CODES, c["caps"])):
        x = ox + QRX + i * (QRSIZE + QRGAP)
        slide.shapes.add_picture(str(png), Mm(x), Mm(oy + 25),
                                 Mm(QRSIZE), Mm(QRSIZE))
        text(slide, x - 3, oy + 69, QRSIZE + 6, 6, cap,
             font="Trebuchet MS", size=10, colour=MAROON, bold=True,
             spacing=1.5, align=PP_ALIGN.CENTER)


for url, png in CODES:
    # 33 modules at 36px: the same 1188px image the review code always had.
    segno.make(url, error="m").save(str(png), kind="png", scale=36, border=2,
                                    dark="#531C23", light="#FBF7F0")

prs = Presentation()
prs.slide_width, prs.slide_height = Mm(PAGE_W), Mm(PAGE_H)
# The dimensions alone leave PowerPoint reporting "Custom". The sldSz type
# attribute is what makes the Slide Size dialog say A4, which is what the
# person printing it needs to see.
prs._element.find(qn("p:sldSz")).set("type", "A4")  # the enum is capitalised
blank = prs.slide_layouts[6]

slide = prs.slides.add_slide(blank)
bg = slide.background.fill
bg.solid()
bg.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for oy, c in zip(STRIP_Y, CONTENT):
    strip(slide, oy, c)

# Where to cut.
line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                  Mm(SHEET), Mm(CUT_Y), Mm(SHEET + W), Mm(CUT_Y))
line.line.color.rgb = RGBColor(0xC4, 0xBC, 0xB0)
line.line.width = Pt(0.75)
line.line.dash_style = MSO_LINE_DASH_STYLE.DASH

slide.notes_slide.notes_text_frame.text = NOTES

# BANNER_PPTX lets a check run write somewhere else, so the real file is not
# clobbered while it is open in PowerPoint.
out = pathlib.Path(os.environ.get("BANNER_PPTX") or ROOT / "print" / "banner.pptx")
prs.save(out)
print("wrote {}  ({} KB)".format(out, out.stat().st_size // 1024))
